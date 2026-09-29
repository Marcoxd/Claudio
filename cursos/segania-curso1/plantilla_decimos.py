"""Plantilla Excel de cálculo de décimos (recurso de la lección 3.3)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, BLUE, SOFT, INPUT = "0D2640", "1D4F8C", "E9F0F9", "FFF6E0"
thin = Side(style="thin", color="D5DDE8")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

# ---------- Décimo tercero ----------
ws = wb.active
ws.title = "Décimo tercero"
ws["A1"] = "Cálculo del décimo tercero"
ws["A1"].font = Font(bold=True, size=16, color=NAVY)
ws["A2"] = "Período: 1 de diciembre al 30 de noviembre. Llena las celdas amarillas."
ws["A2"].font = Font(italic=True, color="5A6A7B")
hdr = ["Mes", "Sueldo", "Horas extra", "Comisiones", "Bonificaciones", "Total del mes", "Décimo tercero mensual"]
for c, h in enumerate(hdr, 1):
    cell = ws.cell(row=4, column=c, value=h)
    cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); cell.border = box
meses = ["Diciembre", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre"]
for i, m in enumerate(meses):
    r = 5 + i
    ws.cell(row=r, column=1, value=m).border = box
    for c in range(2, 6):
        cell = ws.cell(row=r, column=c, value=600 if c == 2 else 0)
        cell.fill = PatternFill("solid", fgColor=INPUT); cell.number_format = '"$"#,##0.00'; cell.border = box
    ws.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})").number_format = '"$"#,##0.00'
    ws.cell(row=r, column=7, value=f"=F{r}/12").number_format = '"$"#,##0.00'
    ws.cell(row=r, column=6).border = box; ws.cell(row=r, column=7).border = box
ws["A17"] = "TOTAL"; ws["A17"].font = Font(bold=True)
for c in "BCDEFG":
    ws[f"{c}17"] = f"=SUM({c}5:{c}16)"; ws[f"{c}17"].font = Font(bold=True); ws[f"{c}17"].number_format = '"$"#,##0.00'
    ws[f"{c}17"].fill = PatternFill("solid", fgColor=SOFT); ws[f"{c}17"].border = box
ws["A19"] = "Décimo tercero acumulado (pago hasta el 24 de diciembre)"; ws["A19"].font = Font(bold=True, color=NAVY)
ws["G19"] = "=F17/12"; ws["G19"].number_format = '"$"#,##0.00'; ws["G19"].font = Font(bold=True, size=14, color=BLUE)
ws["A21"] = "No incluir: utilidades, viáticos, décimo cuarto ni fondos de reserva (Art. 95 y 111 del Código del Trabajo)."
ws["A21"].font = Font(italic=True, color="5A6A7B")
for col, w in zip("ABCDEFG", [16, 13, 13, 13, 15, 15, 22]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[4].height = 32

# ---------- Décimo cuarto ----------
w2 = wb.create_sheet("Décimo cuarto")
w2["A1"] = "Cálculo del décimo cuarto"; w2["A1"].font = Font(bold=True, size=16, color=NAVY)
w2["A2"] = "Llena las celdas amarillas."; w2["A2"].font = Font(italic=True, color="5A6A7B")
rows = [
    ("Sueldo básico unificado vigente", 482, '"$"#,##0.00', True),
    ("Región del lugar de trabajo", "Sierra y Amazonía", "@", True),
    ("Meses trabajados en el período (0 a 12)", 6, "0", True),
    ("Horas semanales de la jornada (40 = completa)", 40, "0", True),
    ("Fecha límite de pago (acumulado)", '=IF(B5="Costa y Galápagos","15 de marzo","15 de agosto")', "@", False),
    ("Período de cálculo", '=IF(B5="Costa y Galápagos","1 de marzo al último día de febrero","1 de agosto al 31 de julio")', "@", False),
    ("Décimo cuarto acumulado", "=ROUND(B4/12*B6*MIN(B7,40)/40,2)", '"$"#,##0.00', False),
    ("Décimo cuarto mensual (jornada completa)", "=ROUND(B4/12*MIN(B7,40)/40,2)", '"$"#,##0.00', False),
]
for i, (lab, val, fmt, editable) in enumerate(rows):
    r = 4 + i
    a = w2.cell(row=r, column=1, value=lab); a.border = box
    b = w2.cell(row=r, column=2, value=val); b.number_format = fmt; b.border = box
    b.fill = PatternFill("solid", fgColor=INPUT if editable else SOFT)
    if not editable:
        b.font = Font(bold=True, color=BLUE)
w2["B10"].font = Font(bold=True, size=14, color=BLUE)
dv = DataValidation(type="list", formula1='"Costa y Galápagos,Sierra y Amazonía"', allow_blank=False)
w2.add_data_validation(dv); dv.add("B5")
w2["A13"] = "Base legal: Art. 113 del Código del Trabajo. SBU 2026: USD 482 (Acuerdo MDT-2025-195)."
w2["A13"].font = Font(italic=True, color="5A6A7B")
w2.column_dimensions["A"].width = 48; w2.column_dimensions["B"].width = 34

import os
os.makedirs("salida/3.3", exist_ok=True)
wb.save("salida/3.3/plantilla-calculo-decimos.xlsx")
print("ok")
